import datetime
import io

import openpyxl
import pytest
from django.core.exceptions import ValidationError
from django.urls import reverse

from apps.catalog.models import Product
from apps.imports import services
from apps.inventory.models import InventoryTransaction, StockBalance, UnitAsset

from .test_imports_services import _base_row, _csv_upload

pytestmark = pytest.mark.django_db


def batch_for(user, **values):
    raw = _base_row(
        **{
            "LOCATION": "",
            "Source Room": "Room A",
            "Status": "In Use",
            "Arrival Date": "2026-01-01",
            "Movement Date": "2026-01-02",
            "Installation Notes": "Installed in server room X",
            "Tracking Method": "unit",
            **values,
        }
    )
    return services.create_batch_from_upload(uploaded_file=_csv_upload([raw]), user=user)[0]


@pytest.mark.parametrize("status", ["Assigned", "Delivered", "In Use"])
def test_issued_unit_creates_receipt_and_movement_once(administrator, location_tree, status):
    batch = batch_for(
        administrator,
        **{
            "Status": status,
            "Employee": "Example Employee",
            "FINAL CUSTOMER": "Customer",
            "Project Ref. #": "P1",
            "S/N": "",
        },
    )
    services.execute_batch(batch=batch, user=administrator)
    row = batch.rows.get()
    assert row.outcome == "imported", row.outcome_detail
    asset = row.created_unit_asset
    assert asset.status == status.lower().replace(" ", "_")
    assert asset.current_location_id is None
    assert InventoryTransaction.objects.count() == 2
    services.execute_batch(batch=batch, user=administrator)
    assert InventoryTransaction.objects.count() == 2


@pytest.mark.parametrize("status", ["Assigned", "Delivered"])
def test_quantity_issue_leaves_no_available_balance(administrator, location_tree, status):
    batch = batch_for(
        administrator,
        **{
            "Status": status,
            "Employee": "Employee",
            "FINAL CUSTOMER": "Customer",
            "Project Ref. #": "P1",
            "S/N": "",
            "Tracking Method": "quantity",
            "QTY": "12",
        },
    )
    services.execute_batch(batch=batch, user=administrator)
    row = batch.rows.get()
    assert row.outcome == "imported", row.outcome_detail
    assert StockBalance.objects.get().on_hand_quantity == 0
    assert not UnitAsset.objects.exists()


def test_blank_location_never_defaults_to_available_stock(administrator, location_tree):
    batch, _ = services.create_batch_from_upload(
        user=administrator,
        uploaded_file=_csv_upload([_base_row(LOCATION="")]),
        default_location=location_tree["room"],
    )
    services.execute_batch(batch=batch, user=administrator)
    assert batch.rows.get().outcome == "warning"
    assert not InventoryTransaction.objects.exists()


@pytest.mark.parametrize(
    "changes",
    [
        {"Employee": "", "Status": "Assigned"},
        {"Status": "Delivered", "FINAL CUSTOMER": ""},
        {"Installation Notes": ""},
        {"Movement Date": "invalid"},
        {"Movement Date": "2025-01-01"},
        {"Tracking Method": "quantity", "S/N": ""},
        {"Stock Purpose": "Customer"},
        {"Status": "Reserved"},
    ],
)
def test_incomplete_lifecycle_cannot_create_stock(administrator, location_tree, changes):
    batch = batch_for(administrator, **changes)
    services.execute_batch(batch=batch, user=administrator)
    assert batch.rows.get().outcome == "warning"
    assert not InventoryTransaction.objects.exists()


def test_movement_failure_rolls_back_receipt_and_catalog(administrator, location_tree, monkeypatch):
    batch = batch_for(administrator)

    def fail(**kwargs):
        raise ValidationError("Movement rejected")

    monkeypatch.setattr(services, "apply_import_movement", fail)
    services.execute_batch(batch=batch, user=administrator)
    assert batch.rows.get().outcome == "failed"
    assert not InventoryTransaction.objects.exists()
    assert not UnitAsset.objects.exists()
    assert not Product.objects.exists()


def test_status_review_preserves_source_and_can_execute(client, administrator, location_tree):
    batch = batch_for(administrator, Status="", **{"Source Room": ""})
    row = batch.rows.get()
    original = row.raw_data.copy()
    client.force_login(administrator)
    url = reverse("imports:row_lifecycle", args=[batch.pk, row.pk])
    assert client.get(url).status_code == 200
    response = client.post(
        url,
        {
            "import_status": "in_use",
            "location": location_tree["room"].pk,
            "arrival_date": "2026-01-01",
            "movement_date": "2026-01-02",
            "installation_notes": "Room X",
        },
    )
    assert response.status_code == 302
    services.execute_batch(batch=batch, user=administrator)
    row.refresh_from_db()
    assert row.raw_data == original
    assert row.created_unit_asset.status == "in_use"
    with pytest.raises(ValidationError):
        services.review_row_lifecycle(
            row=row,
            user=administrator,
            location=location_tree["room"],
            import_status="in_stock",
            arrival_date=datetime.date(2026, 1, 1),
        )


def test_manager_cannot_review_import(client, stock_manager, administrator, location_tree):
    batch = batch_for(administrator)
    client.force_login(stock_manager)
    assert (
        client.get(
            reverse("imports:row_lifecycle", args=[batch.pk, batch.rows.get().pk])
        ).status_code
        == 403
    )


def test_unit_quantity_cannot_silently_drop_items(location_tree):
    _, outcome, detail = services._stage_row(
        _base_row(**{"Tracking Method": "unit", "QTY": "3", "S/N": ""})
    )
    assert outcome == "failed"
    assert "one row per item" in detail


def test_excel_template_has_lifecycle_choices_and_instructions():
    workbook = openpyxl.load_workbook(io.BytesIO(services.build_template_xlsx()))
    assert "Instructions" in workbook.sheetnames
    sheet = workbook.active
    assert sheet.freeze_panes == "A2"
    assert len(sheet.data_validations.dataValidation) == 3
    assert "In Use" in sheet.data_validations.dataValidation[0].formula1
    assert sheet.max_column == len(services.parsing.COLUMNS)
    assert sheet.cell(1, 7).value == "Shelf/Rack"
    assert (
        services.parsing._map_headers([cell.value for cell in sheet[1]])["2nd floor Location"] == 6
    )


def test_imported_row_cannot_be_edited_in_partial_batch(administrator, location_tree):
    batch, _ = services.create_batch_from_upload(
        user=administrator,
        uploaded_file=_csv_upload([_base_row(), _base_row(LOCATION="", **{"S/N": "SECOND"})]),
    )
    services.execute_batch(batch=batch, user=administrator)
    row = batch.rows.get(outcome="imported")
    for mutation in (services.skip_row, services.acknowledge_row_duplicate_serial):
        with pytest.raises(ValidationError):
            mutation(row=row, user=administrator)
    with pytest.raises(ValidationError):
        services.set_row_location_override(
            row=row, user=administrator, location=location_tree["room"]
        )
    row.refresh_from_db()
    assert row.outcome == "imported"


def test_source_override_rejects_country(administrator, location_tree):
    batch = batch_for(administrator)
    with pytest.raises(ValidationError):
        services.set_row_location_override(
            row=batch.rows.get(), user=administrator, location=location_tree["country"]
        )


def test_execution_message_uses_actual_import_count(client, administrator, location_tree):
    batch = batch_for(administrator)
    client.force_login(administrator)
    response = client.post(reverse("imports:execute", args=[batch.pk]), follow=True)
    assert b"Import finished: 1 imported" in response.content
